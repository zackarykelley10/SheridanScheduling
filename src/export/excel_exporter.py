import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from typing import Dict, List, Any
from ..models import Student, Teacher, Course, Room, TimeSlot, DEFAULT_SCHEDULE


class ExcelExporter:
    """Excel export system for generating formatted schedules."""
    
    def __init__(self, students: Dict[str, Student], teachers: Dict[str, Teacher],
                 courses: Dict[str, Course], rooms: Dict[str, Room]):
        self.students = students
        self.teachers = teachers
        self.courses = courses
        self.rooms = rooms
        self.time_slots = DEFAULT_SCHEDULE
    
    def export_all_schedules(self, output_file: str) -> None:
        """Export all schedules to a single Excel file with multiple sheets."""
        wb = Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Create individual student schedules
        self._create_student_schedules_sheet(wb)
        
        # Create master schedule
        self._create_master_schedule_sheet(wb)
        
        # Create teacher schedules
        self._create_teacher_schedules_sheet(wb)
        
        # Create room utilization report
        self._create_room_utilization_sheet(wb)
        
        # Create summary report
        self._create_summary_sheet(wb)
        
        # Save the workbook
        wb.save(output_file)
    
    def _create_student_schedules_sheet(self, wb: Workbook) -> None:
        """Create individual student schedule sheet."""
        ws = wb.create_sheet("Student Schedules")
        
        # Headers
        headers = ["Student ID", "Student Name", "Grade"] + [f"Period {i}" for i in range(1, 8)]
        ws.append(headers)
        
        # Format headers
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Add student data
        for student in self.students.values():
            row_data = [student.id, student.name, student.grade]
            
            # Create schedule array for 7 periods
            schedule = [""] * 7
            
            for course_id in student.assigned_classes:
                if course_id in self.courses:
                    course = self.courses[course_id]
                    if course.assigned_period is not None:
                        period_index = course.assigned_period - 1
                        teacher_name = "TBD"
                        room_name = "TBD"
                        
                        if course.assigned_teacher and course.assigned_teacher in self.teachers:
                            teacher_name = self.teachers[course.assigned_teacher].name
                        
                        if course.assigned_room and course.assigned_room in self.rooms:
                            room_name = self.rooms[course.assigned_room].name
                        
                        schedule[period_index] = f"{course.name}\n{teacher_name}\nRoom: {room_name}"
            
            row_data.extend(schedule)
            ws.append(row_data)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 20)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_master_schedule_sheet(self, wb: Workbook) -> None:
        """Create master schedule showing all class assignments."""
        ws = wb.create_sheet("Master Schedule")
        
        # Headers
        headers = ["Course ID", "Course Name", "Subject", "Teacher", "Room", "Period", 
                  "Time", "Enrolled", "Capacity", "Utilization %"]
        ws.append(headers)
        
        # Format headers
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Add course data
        for course in self.courses.values():
            teacher_name = "Unassigned"
            if course.assigned_teacher and course.assigned_teacher in self.teachers:
                teacher_name = self.teachers[course.assigned_teacher].name
            
            room_name = "Unassigned"
            if course.assigned_room and course.assigned_room in self.rooms:
                room_name = self.rooms[course.assigned_room].name
            
            period_info = "Unassigned"
            time_info = "TBD"
            if course.assigned_period is not None:
                period_info = str(course.assigned_period)
                if 1 <= course.assigned_period <= 7:
                    time_slot = self.time_slots[course.assigned_period - 1]
                    time_info = f"{time_slot.start_time} - {time_slot.end_time}"
            
            enrolled = course.get_enrollment_count()
            utilization = (enrolled / course.capacity * 100) if course.capacity > 0 else 0
            
            row_data = [
                course.id, course.name, course.subject, teacher_name, room_name,
                period_info, time_info, enrolled, course.capacity, f"{utilization:.1f}%"
            ]
            ws.append(row_data)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 15)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_teacher_schedules_sheet(self, wb: Workbook) -> None:
        """Create teacher schedule sheet."""
        ws = wb.create_sheet("Teacher Schedules")
        
        # Headers
        headers = ["Teacher ID", "Teacher Name"] + [f"Period {i}" for i in range(1, 8)]
        ws.append(headers)
        
        # Format headers
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Add teacher data
        for teacher in self.teachers.values():
            row_data = [teacher.id, teacher.name]
            
            # Create schedule array for 7 periods
            schedule = [""] * 7
            
            for course_id in teacher.assigned_classes:
                if course_id in self.courses:
                    course = self.courses[course_id]
                    if course.assigned_period is not None:
                        period_index = course.assigned_period - 1
                        room_name = "TBD"
                        
                        if course.assigned_room and course.assigned_room in self.rooms:
                            room_name = self.rooms[course.assigned_room].name
                        
                        enrolled_count = course.get_enrollment_count()
                        schedule[period_index] = f"{course.name}\nRoom: {room_name}\nStudents: {enrolled_count}"
            
            row_data.extend(schedule)
            ws.append(row_data)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 18)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_room_utilization_sheet(self, wb: Workbook) -> None:
        """Create room utilization report."""
        ws = wb.create_sheet("Room Utilization")
        
        # Headers
        headers = ["Room ID", "Room Name", "Capacity"] + [f"Period {i}" for i in range(1, 8)] + ["Utilization %"]
        ws.append(headers)
        
        # Format headers
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        
        # Add room data
        for room in self.rooms.values():
            row_data = [room.id, room.name, room.capacity]
            
            # Period assignments
            periods_used = 0
            for period in range(1, 8):
                course_id = room.schedule.get(period)
                if course_id and course_id in self.courses:
                    course = self.courses[course_id]
                    enrolled = course.get_enrollment_count()
                    row_data.append(f"{course.name}\n({enrolled} students)")
                    periods_used += 1
                else:
                    row_data.append("Available")
            
            # Calculate utilization percentage
            utilization = (periods_used / 7 * 100)
            row_data.append(f"{utilization:.1f}%")
            
            ws.append(row_data)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 15)
            ws.column_dimensions[column_letter].width = adjusted_width
    
    def _create_summary_sheet(self, wb: Workbook) -> None:
        """Create summary statistics sheet."""
        ws = wb.create_sheet("Summary")
        
        # Calculate statistics
        total_students = len(self.students)
        total_courses = len(self.courses)
        total_teachers = len(self.teachers)
        total_rooms = len(self.rooms)
        
        assigned_courses = len([c for c in self.courses.values() if c.is_assigned()])
        total_enrollments = sum(c.get_enrollment_count() for c in self.courses.values())
        
        # Create summary data
        summary_data = [
            ["SCHEDULING SUMMARY", ""],
            ["", ""],
            ["Total Students", total_students],
            ["Total Courses", total_courses],
            ["Total Teachers", total_teachers],
            ["Total Rooms", total_rooms],
            ["", ""],
            ["Assigned Courses", assigned_courses],
            ["Unassigned Courses", total_courses - assigned_courses],
            ["Total Enrollments", total_enrollments],
            ["", ""],
            ["COURSE UTILIZATION", ""],
        ]
        
        # Add course utilization details
        for course in self.courses.values():
            enrolled = course.get_enrollment_count()
            utilization = (enrolled / course.capacity * 100) if course.capacity > 0 else 0
            summary_data.append([f"{course.name}", f"{enrolled}/{course.capacity} ({utilization:.1f}%)"])
        
        # Write data to sheet
        for row_data in summary_data:
            ws.append(row_data)
        
        # Format the summary sheet
        ws.cell(row=1, column=1).font = Font(size=16, bold=True)
        ws.cell(row=12, column=1).font = Font(size=14, bold=True)
        
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column_letter].width = adjusted_width