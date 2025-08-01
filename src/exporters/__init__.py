"""
Excel exporter for generating formatted schedule reports.
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows
from typing import Dict, List, Any
from pathlib import Path
import logging

from ..models import Schedule, Student, Teacher, ClassOffering, Room, Period

logger = logging.getLogger(__name__)


class ExcelExporter:
    """Export schedules to formatted Excel files with multiple views."""
    
    def __init__(self, schedule: Schedule, students: Dict[str, Student], 
                 teachers: Dict[str, Teacher], classes: Dict[str, ClassOffering], 
                 rooms: Dict[str, Room]):
        self.schedule = schedule
        self.students = students
        self.teachers = teachers
        self.classes = classes
        self.rooms = rooms
    
    def export_all_views(self, output_path: str) -> None:
        """Export all schedule views to a single Excel file."""
        logger.info(f"Exporting all schedule views to {output_path}")
        
        wb = Workbook()
        
        # Remove default sheet
        wb.remove(wb.active)
        
        # Create all views
        self._create_master_schedule(wb)
        self._create_student_schedules(wb)
        self._create_teacher_schedules(wb)
        self._create_room_schedules(wb)
        self._create_summary_sheet(wb)
        
        # Save workbook
        wb.save(output_path)
        logger.info(f"Excel export completed: {output_path}")
    
    def _create_master_schedule(self, wb: Workbook) -> None:
        """Create master schedule showing all classes by period."""
        ws = wb.create_sheet("Master Schedule")
        
        # Create header
        headers = ["Period", "Class", "Teacher", "Room", "Students", "Capacity"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Add data
        row = 2
        for period in Period:
            period_entries = self.schedule.get_period_schedule(period)
            if not period_entries:
                continue
                
            for entry in period_entries:
                class_name = self.classes[entry.class_id].name
                teacher_name = self.teachers[entry.teacher_id].name
                room_name = self.rooms[entry.room_id].name
                student_count = len(entry.students)
                max_capacity = self.classes[entry.class_id].max_capacity
                
                ws.cell(row=row, column=1, value=f"Period {period.value}")
                ws.cell(row=row, column=2, value=class_name)
                ws.cell(row=row, column=3, value=teacher_name)
                ws.cell(row=row, column=4, value=room_name)
                ws.cell(row=row, column=5, value=student_count)
                ws.cell(row=row, column=6, value=f"{student_count}/{max_capacity}")
                row += 1
        
        self._format_worksheet(ws)
    
    def _create_student_schedules(self, wb: Workbook) -> None:
        """Create individual student schedules."""
        ws = wb.create_sheet("Student Schedules")
        
        # Create header
        headers = ["Student ID", "Student Name", "Grade"] + [f"Period {i}" for i in range(1, 8)]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Add student data
        row = 2
        for student in self.students.values():
            student_schedule = self.schedule.get_student_schedule(student.id)
            
            # Create period mapping
            period_classes = {}
            for entry in student_schedule:
                class_name = self.classes[entry.class_id].name
                room_name = self.rooms[entry.room_id].name
                period_classes[entry.period.value] = f"{class_name} (Rm {room_name})"
            
            ws.cell(row=row, column=1, value=student.id)
            ws.cell(row=row, column=2, value=student.name)
            ws.cell(row=row, column=3, value=student.grade)
            
            # Fill period columns
            for period in range(1, 8):
                col = 3 + period
                class_info = period_classes.get(period, "Free Period")
                ws.cell(row=row, column=col, value=class_info)
            
            row += 1
        
        self._format_worksheet(ws)
    
    def _create_teacher_schedules(self, wb: Workbook) -> None:
        """Create teacher schedule view."""
        ws = wb.create_sheet("Teacher Schedules")
        
        # Create header
        headers = ["Teacher ID", "Teacher Name"] + [f"Period {i}" for i in range(1, 8)]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Add teacher data
        row = 2
        for teacher in self.teachers.values():
            teacher_schedule = self.schedule.get_teacher_schedule(teacher.id)
            
            # Create period mapping
            period_classes = {}
            for entry in teacher_schedule:
                class_name = self.classes[entry.class_id].name
                room_name = self.rooms[entry.room_id].name
                student_count = len(entry.students)
                period_classes[entry.period.value] = f"{class_name} (Rm {room_name}) [{student_count} students]"
            
            ws.cell(row=row, column=1, value=teacher.id)
            ws.cell(row=row, column=2, value=teacher.name)
            
            # Fill period columns
            for period in range(1, 8):
                col = 2 + period
                class_info = period_classes.get(period, "Free Period")
                ws.cell(row=row, column=col, value=class_info)
            
            row += 1
        
        self._format_worksheet(ws)
    
    def _create_room_schedules(self, wb: Workbook) -> None:
        """Create room usage schedule."""
        ws = wb.create_sheet("Room Schedules")
        
        # Create header
        headers = ["Room ID", "Room Name", "Capacity"] + [f"Period {i}" for i in range(1, 8)]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF")
        
        # Add room data
        row = 2
        for room in self.rooms.values():
            room_schedule = self.schedule.get_room_schedule(room.id)
            
            # Create period mapping
            period_classes = {}
            for entry in room_schedule:
                class_name = self.classes[entry.class_id].name
                teacher_name = self.teachers[entry.teacher_id].name
                student_count = len(entry.students)
                period_classes[entry.period.value] = f"{class_name} - {teacher_name} [{student_count} students]"
            
            ws.cell(row=row, column=1, value=room.id)
            ws.cell(row=row, column=2, value=room.name)
            ws.cell(row=row, column=3, value=room.capacity)
            
            # Fill period columns
            for period in range(1, 8):
                col = 3 + period
                class_info = period_classes.get(period, "Available")
                ws.cell(row=row, column=col, value=class_info)
            
            row += 1
        
        self._format_worksheet(ws)
    
    def _create_summary_sheet(self, wb: Workbook) -> None:
        """Create summary statistics sheet."""
        ws = wb.create_sheet("Summary", 0)  # Insert as first sheet
        
        # Title
        ws.cell(row=1, column=1, value="Sheridan High School - Schedule Summary")
        title_cell = ws.cell(row=1, column=1)
        title_cell.font = Font(size=16, bold=True)
        
        # Statistics
        row = 3
        stats = self._calculate_summary_stats()
        
        for key, value in stats.items():
            ws.cell(row=row, column=1, value=key.replace('_', ' ').title())
            ws.cell(row=row, column=2, value=value)
            row += 1
        
        # Format summary
        for row_cells in ws.iter_rows(min_row=3, max_row=row-1, min_col=1, max_col=2):
            row_cells[0].font = Font(bold=True)
            row_cells[0].fill = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
        
        self._format_worksheet(ws)
    
    def _calculate_summary_stats(self) -> Dict[str, Any]:
        """Calculate summary statistics."""
        total_students = len(self.students)
        total_teachers = len(self.teachers)
        total_rooms = len(self.rooms)
        total_classes = len(self.classes)
        scheduled_entries = len(self.schedule.entries)
        
        # Calculate utilization
        total_periods = 7
        max_possible_classes = min(total_teachers, total_rooms) * total_periods
        utilization_rate = (scheduled_entries / max_possible_classes) * 100 if max_possible_classes > 0 else 0
        
        # Student enrollment stats
        total_enrollments = sum(len(entry.students) for entry in self.schedule.entries)
        avg_class_size = total_enrollments / scheduled_entries if scheduled_entries > 0 else 0
        
        return {
            'total_students': total_students,
            'total_teachers': total_teachers,
            'total_rooms': total_rooms,
            'total_classes_offered': total_classes,
            'scheduled_class_periods': scheduled_entries,
            'total_student_enrollments': total_enrollments,
            'average_class_size': f"{avg_class_size:.1f}",
            'facility_utilization_rate': f"{utilization_rate:.1f}%"
        }
    
    def _format_worksheet(self, ws) -> None:
        """Apply consistent formatting to worksheet."""
        # Auto-adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Add borders
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        for row in ws.iter_rows():
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical='center')
        
        # Freeze top row
        ws.freeze_panes = "A2"